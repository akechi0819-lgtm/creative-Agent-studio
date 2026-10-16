#!/usr/bin/env python3
"""Small, bounded Office/CSV bridge for task-local workspace files."""

from __future__ import annotations

import csv
import datetime as dt
import json
import os
import sys
import zipfile
from pathlib import Path, PurePosixPath

MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_CELL_TEXT = 2_000
MAX_EXTRACTED_TEXT = 200_000
MAX_WORKBOOK_ROWS = 10_000
MAX_WORKBOOK_COLUMNS = 100
MAX_WORKBOOK_SHEETS = 20
MAX_WRITE_ROWS = 20_000
MAX_WRITE_COLUMNS = 100
MAX_CSV_SAMPLE_BYTES = 16_384


def json_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        if isinstance(value, str) and len(value) > MAX_CELL_TEXT:
            return value[:MAX_CELL_TEXT]
        return value
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)[:MAX_CELL_TEXT]


def safe_path(root: Path, relative: str, *, must_exist: bool = True) -> Path:
    if not isinstance(relative, str) or not relative or "\\" in relative or "\x00" in relative:
        raise ValueError("invalid_relative_path")
    parts = PurePosixPath(relative).parts
    if PurePosixPath(relative).is_absolute() or any(part in ("", ".", "..") for part in parts):
        raise ValueError("invalid_relative_path")
    if parts[0] not in ("inputs", "work", "outputs"):
        raise ValueError("invalid_workspace_area")
    current = root
    for index, part in enumerate(parts):
        current = current / part
        if current.is_symlink():
            raise ValueError("workspace_symlink_not_allowed")
        if not current.exists():
            if must_exist or index < len(parts) - 1:
                raise FileNotFoundError("workspace_file_not_found")
            break
        if index < len(parts) - 1 and not current.is_dir():
            raise ValueError("workspace_parent_not_directory")
    if must_exist:
        resolved = current.resolve(strict=True)
    else:
        resolved = current.parent.resolve(strict=True) / current.name
    if resolved == root or root not in resolved.parents:
        raise ValueError("workspace_path_outside_root")
    if must_exist and not resolved.is_file():
        raise ValueError("workspace_path_not_file")
    return resolved


def check_office_archive(file_path: Path) -> None:
    try:
        with zipfile.ZipFile(file_path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("office_archive_has_too_many_entries")
            expanded = sum(info.file_size for info in infos)
            if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                raise ValueError("office_archive_expands_past_limit")
            for info in infos:
                if info.flag_bits & 0x1:
                    raise ValueError("encrypted_office_archive_not_supported")
                if info.compress_size == 0 and info.file_size > 0:
                    raise ValueError("invalid_office_archive")
                if info.compress_size > 0 and info.file_size / info.compress_size > 200:
                    raise ValueError("office_archive_compression_ratio_too_high")
    except zipfile.BadZipFile as error:
        raise ValueError("invalid_office_archive") from error


def read_xlsx(root: Path, request: dict) -> dict:
    from openpyxl import load_workbook

    file_path = safe_path(root, request.get("relativePath"))
    check_office_archive(file_path)
    start_row = request.get("startRow", 1)
    max_rows = request.get("maxRows", 150)
    if not isinstance(start_row, int) or start_row < 1 or start_row > MAX_WORKBOOK_ROWS:
        raise ValueError("invalid_start_row")
    if not isinstance(max_rows, int) or max_rows < 1 or max_rows > 300:
        raise ValueError("invalid_max_rows")
    workbook = load_workbook(file_path, read_only=True, data_only=False, keep_links=False)
    try:
        names = workbook.sheetnames[:MAX_WORKBOOK_SHEETS]
        selected = request.get("sheetName")
        if selected is not None and selected not in workbook.sheetnames:
            raise ValueError("worksheet_not_found")
        sheets = []
        text_budget = MAX_EXTRACTED_TEXT
        budget_exhausted = False
        for worksheet in workbook.worksheets:
            if selected is not None and worksheet.title != selected:
                continue
            if worksheet.title not in names:
                continue
            max_column = min(worksheet.max_column or 1, MAX_WORKBOOK_COLUMNS)
            end_row = min(start_row + max_rows - 1, MAX_WORKBOOK_ROWS)
            rows = []
            for row_number, values in enumerate(
                worksheet.iter_rows(min_row=start_row, max_row=end_row, max_col=max_column, values_only=True),
                start=start_row,
            ):
                cells = []
                for value in values:
                    cell = json_value(value)
                    if isinstance(cell, str):
                        if text_budget <= 0:
                            cell = ""
                            budget_exhausted = True
                        elif len(cell) > text_budget:
                            cell = cell[:text_budget]
                            budget_exhausted = True
                        text_budget -= len(cell)
                    cells.append(cell)
                while cells and cells[-1] is None:
                    cells.pop()
                if any(value is not None for value in cells):
                    rows.append({"row": row_number, "values": cells})
                if budget_exhausted:
                    break
            known_max_row = worksheet.max_row or 0
            sheets.append({
                "name": worksheet.title,
                "rows": rows,
                "columnCount": max_column,
                "rowCount": known_max_row,
                "truncated": budget_exhausted or known_max_row > end_row or (worksheet.max_column or 0) > MAX_WORKBOOK_COLUMNS,
                "nextRow": (rows[-1]["row"] + 1 if rows and budget_exhausted else end_row + 1) if known_max_row > (rows[-1]["row"] if rows else end_row) else None,
            })
            if budget_exhausted:
                break
        return {"kind": "workbook", "sheets": sheets, "sheetCount": len(workbook.sheetnames)}
    finally:
        workbook.close()


def read_docx(root: Path, request: dict) -> dict:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    file_path = safe_path(root, request.get("relativePath"))
    check_office_archive(file_path)
    document = Document(file_path)
    blocks = []
    text_parts = []
    budget = 50_000
    table_index = 0
    truncated = False
    for child in document.element.body.iterchildren():
        if len(blocks) >= 1_000 or budget <= 0:
            truncated = True
            break
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text.strip()
            if text:
                text = text[:budget]
                blocks.append({"type": "paragraph", "text": text})
                text_parts.append(text)
                budget -= len(text)
        elif child.tag.endswith("}tbl"):
            table_index += 1
            table = Table(child, document)
            rows = []
            for row in table.rows[:500]:
                row_values = []
                for cell in row.cells[:MAX_WORKBOOK_COLUMNS]:
                    value = cell.text[: min(MAX_CELL_TEXT, budget)]
                    budget -= len(value)
                    row_values.append(value)
                    if budget <= 0:
                        truncated = True
                        break
                rows.append(row_values)
                if budget <= 0:
                    break
            blocks.append({"type": "table", "index": table_index, "rows": rows})
            text_parts.extend("\t".join(row) for row in rows)
    text = "\n".join(text_parts)
    return {"kind": "document", "format": "docx", "blocks": blocks, "text": text, "truncated": truncated}


def read_csv(root: Path, request: dict) -> dict:
    file_path = safe_path(root, request.get("relativePath"))
    start_row = request.get("startRow", 1)
    max_rows = request.get("maxRows", 150)
    if not isinstance(start_row, int) or start_row < 1 or start_row > MAX_WORKBOOK_ROWS:
        raise ValueError("invalid_start_row")
    if not isinstance(max_rows, int) or max_rows < 1 or max_rows > 300:
        raise ValueError("invalid_max_rows")
    with file_path.open("r", encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(MAX_CSV_SAMPLE_BYTES)
        stream.seek(0)
        delimiter = request.get("delimiter")
        if delimiter in (",", ";", "\t", "|"):
            reader = csv.reader(stream, delimiter=delimiter)
        else:
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|") if sample else csv.excel
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(stream, dialect)
        rows = []
        next_row = None
        text_budget = MAX_EXTRACTED_TEXT
        for row_number, row in enumerate(reader, start=1):
            if row_number < start_row:
                continue
            if len(rows) >= max_rows:
                next_row = row_number
                break
            values = []
            for value in row[:MAX_WORKBOOK_COLUMNS]:
                clipped = value[: min(MAX_CELL_TEXT, text_budget)]
                values.append(clipped)
                text_budget -= len(clipped)
            rows.append({"row": row_number, "values": values})
            if text_budget <= 0:
                next_row = row_number + 1
                break
    return {"kind": "table", "format": "csv", "rows": rows, "nextRow": next_row, "truncated": next_row is not None}


def write_xlsx(root: Path, request: dict) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    file_path = safe_path(root, request.get("relativePath"), must_exist=False)
    if file_path.suffix.lower() != ".xlsx":
        raise ValueError("workbook_output_must_be_xlsx")
    sheets = request.get("sheets")
    if not isinstance(sheets, list) or not sheets or len(sheets) > MAX_WORKBOOK_SHEETS:
        raise ValueError("invalid_workbook_sheets")
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names = set()
    for index, sheet in enumerate(sheets, start=1):
        if not isinstance(sheet, dict) or not isinstance(sheet.get("rows"), list):
            raise ValueError("invalid_workbook_sheet")
        name = str(sheet.get("name") or f"Sheet{index}").strip()[:31] or f"Sheet{index}"
        if name in used_names:
            raise ValueError("worksheet_names_must_be_unique")
        used_names.add(name)
        rows = sheet["rows"]
        if len(rows) > MAX_WRITE_ROWS:
            raise ValueError("workbook_has_too_many_rows")
        worksheet = workbook.create_sheet(name)
        max_column = 0
        for row_index, row in enumerate(rows, start=1):
            if not isinstance(row, list) or len(row) > MAX_WRITE_COLUMNS:
                raise ValueError("invalid_workbook_row")
            max_column = max(max_column, len(row))
            for column_index, value in enumerate(row, start=1):
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    raise ValueError("workbook_cells_must_be_scalar_values")
                cell = worksheet.cell(row=row_index, column=column_index, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
            if row_index == 1:
                for cell in worksheet[row_index]:
                    cell.font = Font(bold=True)
        worksheet.freeze_panes = "A2"
        if rows and max_column:
            worksheet.auto_filter.ref = f"A1:{get_column_letter(max_column)}{len(rows)}"
    file_path.parent.mkdir(parents=True, exist_ok=True)
    if file_path.exists() or file_path.is_symlink():
        raise FileExistsError("workspace_output_exists")
    workbook.save(file_path)
    workbook.close()
    return {"kind": "workbook", "relativePath": request["relativePath"], "sheetCount": len(sheets), "rowCount": sum(len(sheet["rows"]) for sheet in sheets)}


def main() -> int:
    try:
        request = json.load(sys.stdin)
        if not isinstance(request, dict):
            raise ValueError("invalid_request")
        root_value = request.get("workspaceRoot")
        if not isinstance(root_value, str) or not os.path.isabs(root_value):
            raise ValueError("invalid_workspace_root")
        root = Path(root_value).resolve(strict=True)
        operation = request.get("operation")
        if operation == "read_xlsx":
            result = read_xlsx(root, request)
        elif operation == "read_docx":
            result = read_docx(root, request)
        elif operation == "read_csv":
            result = read_csv(root, request)
        elif operation == "write_xlsx":
            result = write_xlsx(root, request)
        else:
            raise ValueError("unsupported_operation")
        sys.stdout.write(json.dumps({"ok": True, "result": result}, ensure_ascii=False, separators=(",", ":")))
        return 0
    except Exception as error:
        code = str(error) if isinstance(error, (ValueError, FileNotFoundError, FileExistsError)) else "office_processing_failed"
        sys.stdout.write(json.dumps({"ok": False, "error": code}, ensure_ascii=False, separators=(",", ":")))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
